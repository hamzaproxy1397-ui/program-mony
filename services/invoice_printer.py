# -*- coding: utf-8 -*-
"""
نظام طباعة وتصدير الفواتير - دعم الطباعة الحرارية و A4 و PDF و Excel
"""

import os
import sys
from datetime import datetime
from pathlib import Path
import logging

# استيراد مكتبات التصدير
try:
    from reportlab.lib.pagesizes import A4, letter
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch, cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class InvoicePrinter:
    """نظام طباعة وتصدير الفواتير"""

    def __init__(self):
        self.logger = logging.getLogger(__name__)

        # إعداد مسارات الملفات
        self.project_root = Path(__file__).parent.parent
        self.reports_dir = self.project_root / "reports" / "generated"
        self.fonts_dir = self.project_root / "assets" / "fonts"

        # إنشاء المجلدات إذا لم تكن موجودة
        self.reports_dir.mkdir(parents=True, exist_ok=True)

        # إعداد الخطوط العربية
        self.setup_arabic_fonts()

    def setup_arabic_fonts(self):
        """إعداد الخطوط العربية للطباعة"""
        if not REPORTLAB_AVAILABLE:
            return

        try:
            # البحث عن خطوط عربية في النظام
            font_paths = [
                "C:/Windows/Fonts/arial.ttf",
                "C:/Windows/Fonts/tahoma.ttf",
                "/System/Library/Fonts/Arial.ttf",  # macOS
                "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",  # Linux
            ]

            for font_path in font_paths:
                if os.path.exists(font_path):
                    try:
                        pdfmetrics.registerFont(TTFont('Arabic', font_path))
                        self.arabic_font = 'Arabic'
                        break
                    except:
                        continue
            else:
                # استخدام خط افتراضي
                self.arabic_font = 'Helvetica'

        except Exception as e:
            self.logger.warning(f"تعذر تحميل الخطوط العربية: {str(e)}")
            self.arabic_font = 'Helvetica'

    def print_invoice_pdf(self, invoice_data, items_data, output_path=None):
        """طباعة الفاتورة إلى PDF"""
        if not REPORTLAB_AVAILABLE:
            raise ImportError("مكتبة reportlab غير مثبتة. يرجى تثبيتها باستخدام: pip install reportlab")

        try:
            # تحديد مسار الملف
            if output_path is None:
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                filename = f"invoice_{invoice_data['invoice_number']}_{timestamp}.pdf"
                output_path = self.reports_dir / filename

            # إنشاء مستند PDF
            doc = SimpleDocTemplate(
                str(output_path),
                pagesize=A4,
                rightMargin=2*cm,
                leftMargin=2*cm,
                topMargin=2*cm,
                bottomMargin=2*cm
            )

            # إعداد الأنماط
            styles = getSampleStyleSheet()

            # نمط العنوان العربي
            title_style = ParagraphStyle(
                'ArabicTitle',
                parent=styles['Title'],
                fontName=self.arabic_font,
                fontSize=18,
                alignment=2,  # محاذاة يمين
                spaceAfter=20
            )

            # نمط النص العربي
            arabic_style = ParagraphStyle(
                'Arabic',
                parent=styles['Normal'],
                fontName=self.arabic_font,
                fontSize=12,
                alignment=2,  # محاذاة يمين
                spaceAfter=10
            )

            # بناء محتوى PDF
            story = []

            # عنوان الشركة
            story.append(Paragraph("برنامج ست الكل للمحاسبة", title_style))
            story.append(Paragraph("فاتورة بيع", arabic_style))
            story.append(Spacer(1, 20))

            # معلومات الفاتورة
            invoice_info = [
                [f"رقم الفاتورة: {invoice_data['invoice_number']}", f"التاريخ: {invoice_data['invoice_date']}"],
                [f"العميل: {invoice_data.get('customer_name', 'عميل نقدي')}", f"طريقة الدفع: {invoice_data.get('payment_method', 'نقدي')}"]
            ]

            info_table = Table(invoice_info, colWidths=[8*cm, 8*cm])
            info_table.setStyle(TableStyle([
                ('FONTNAME', (0, 0), (-1, -1), self.arabic_font),
                ('FONTSIZE', (0, 0), (-1, -1), 12),
                ('ALIGN', (0, 0), (-1, -1), 'RIGHT'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('BACKGROUND', (0, 0), (-1, -1), colors.lightgrey)
            ]))

            story.append(info_table)
            story.append(Spacer(1, 20))

            # جدول الأصناف
            items_header = ['الإجمالي', 'السعر', 'الكمية', 'اسم الصنف', 'كود الصنف']
            items_table_data = [items_header]

            total_amount = 0
            for item in items_data:
                row = [
                    f"{item.get('line_total', 0):.2f}",
                    f"{item.get('unit_price', 0):.2f}",
                    f"{item.get('quantity', 0):.2f}",
                    item.get('product_name', ''),
                    item.get('product_code', '')
                ]
                items_table_data.append(row)
                total_amount += item.get('line_total', 0)

            items_table = Table(items_table_data, colWidths=[3*cm, 3*cm, 2*cm, 6*cm, 3*cm])
            items_table.setStyle(TableStyle([
                ('FONTNAME', (0, 0), (-1, -1), self.arabic_font),
                ('FONTSIZE', (0, 0), (-1, -1), 10),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('FONTSIZE', (0, 0), (-1, 0), 12),
                ('FONTNAME', (0, 0), (-1, 0), self.arabic_font)
            ]))

            story.append(items_table)
            story.append(Spacer(1, 20))

            # الإجماليات
            totals_data = [
                ['المجموع الفرعي:', f"{invoice_data.get('subtotal', total_amount):.2f} ر.س"],
                ['الخصم:', f"{invoice_data.get('discount_amount', 0):.2f} ر.س"],
                ['الضريبة:', f"{invoice_data.get('tax_amount', 0):.2f} ر.س"],
                ['الإجمالي النهائي:', f"{invoice_data.get('total_amount', total_amount):.2f} ر.س"]
            ]

            totals_table = Table(totals_data, colWidths=[4*cm, 4*cm])
            totals_table.setStyle(TableStyle([
                ('FONTNAME', (0, 0), (-1, -1), self.arabic_font),
                ('FONTSIZE', (0, 0), (-1, -1), 12),
                ('ALIGN', (0, 0), (-1, -1), 'RIGHT'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('BACKGROUND', (0, -1), (-1, -1), colors.lightblue),
                ('FONTSIZE', (0, -1), (-1, -1), 14),
                ('FONTNAME', (0, -1), (-1, -1), self.arabic_font)
            ]))

            story.append(totals_table)
            story.append(Spacer(1, 30))

            # ملاحظات
            if invoice_data.get('notes'):
                story.append(Paragraph(f"ملاحظات: {invoice_data['notes']}", arabic_style))
                story.append(Spacer(1, 20))

            # توقيع
            signature_data = [
                ['توقيع العميل', 'توقيع المحاسب'],
                ['', '']
            ]

            signature_table = Table(signature_data, colWidths=[8*cm, 8*cm], rowHeights=[1*cm, 2*cm])
            signature_table.setStyle(TableStyle([
                ('FONTNAME', (0, 0), (-1, -1), self.arabic_font),
                ('FONTSIZE', (0, 0), (-1, -1), 12),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))

            story.append(signature_table)

            # بناء المستند
            doc.build(story)

            self.logger.info(f"تم إنشاء ملف PDF: {output_path}")
            return str(output_path)

        except Exception as e:
            self.logger.error(f"خطأ في إنشاء PDF: {str(e)}")
            raise

    def export_invoice_excel(self, invoice_data, items_data, output_path=None):
        """تصدير الفاتورة إلى Excel"""
        if not OPENPYXL_AVAILABLE:
            raise ImportError("مكتبة openpyxl غير مثبتة. يرجى تثبيتها باستخدام: pip install openpyxl")

        try:
            # تحديد مسار الملف
            if output_path is None:
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                filename = f"invoice_{invoice_data['invoice_number']}_{timestamp}.xlsx"
                output_path = self.reports_dir / filename

            # إنشاء مصنف Excel
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "فاتورة"

            # إعداد الأنماط
            header_font = Font(name='Arial', size=14, bold=True)
            normal_font = Font(name='Arial', size=12)
            border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

            # عنوان الفاتورة
            ws['A1'] = "برنامج ست الكل للمحاسبة"
            ws['A1'].font = header_font
            ws['A1'].alignment = Alignment(horizontal='center')
            ws.merge_cells('A1:E1')

            ws['A2'] = "فاتورة بيع"
            ws['A2'].font = header_font
            ws['A2'].alignment = Alignment(horizontal='center')
            ws.merge_cells('A2:E2')

            # معلومات الفاتورة
            row = 4
            ws[f'A{row}'] = "رقم الفاتورة:"
            ws[f'B{row}'] = invoice_data['invoice_number']
            ws[f'D{row}'] = "التاريخ:"
            ws[f'E{row}'] = invoice_data['invoice_date']

            row += 1
            ws[f'A{row}'] = "العميل:"
            ws[f'B{row}'] = invoice_data.get('customer_name', 'عميل نقدي')
            ws[f'D{row}'] = "طريقة الدفع:"
            ws[f'E{row}'] = invoice_data.get('payment_method', 'نقدي')

            # جدول الأصناف
            row += 3
            headers = ['كود الصنف', 'اسم الصنف', 'الكمية', 'السعر', 'الإجمالي']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col, value=header)
                cell.font = header_font
                cell.border = border
                cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
                cell.alignment = Alignment(horizontal='center')

            # بيانات الأصناف
            total_amount = 0
            for item in items_data:
                row += 1
                ws.cell(row=row, column=1, value=item.get('product_code', '')).border = border
                ws.cell(row=row, column=2, value=item.get('product_name', '')).border = border
                ws.cell(row=row, column=3, value=item.get('quantity', 0)).border = border
                ws.cell(row=row, column=4, value=item.get('unit_price', 0)).border = border
                ws.cell(row=row, column=5, value=item.get('line_total', 0)).border = border
                total_amount += item.get('line_total', 0)

            # الإجماليات
            row += 2
            ws[f'D{row}'] = "المجموع الفرعي:"
            ws[f'E{row}'] = invoice_data.get('subtotal', total_amount)

            row += 1
            ws[f'D{row}'] = "الخصم:"
            ws[f'E{row}'] = invoice_data.get('discount_amount', 0)

            row += 1
            ws[f'D{row}'] = "الضريبة:"
            ws[f'E{row}'] = invoice_data.get('tax_amount', 0)

            row += 1
            ws[f'D{row}'] = "الإجمالي النهائي:"
            ws[f'E{row}'] = invoice_data.get('total_amount', total_amount)
            ws[f'D{row}'].font = header_font
            ws[f'E{row}'].font = header_font

            # تنسيق الأعمدة
            ws.column_dimensions['A'].width = 15
            ws.column_dimensions['B'].width = 30
            ws.column_dimensions['C'].width = 10
            ws.column_dimensions['D'].width = 15
            ws.column_dimensions['E'].width = 15

            # حفظ الملف
            wb.save(output_path)

            self.logger.info(f"تم إنشاء ملف Excel: {output_path}")
            return str(output_path)

        except Exception as e:
            self.logger.error(f"خطأ في إنشاء Excel: {str(e)}")
            raise

    def print_thermal_receipt(self, invoice_data, items_data):
        """طباعة إيصال حراري (للطابعات الحرارية)"""
        try:
            # إنشاء نص الإيصال
            receipt_text = self.generate_thermal_receipt_text(invoice_data, items_data)

            # حفظ في ملف نصي للمعاينة
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"thermal_receipt_{invoice_data['invoice_number']}_{timestamp}.txt"
            output_path = self.reports_dir / filename

            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(receipt_text)

            self.logger.info(f"تم إنشاء إيصال حراري: {output_path}")
            return str(output_path)

        except Exception as e:
            self.logger.error(f"خطأ في إنشاء الإيصال الحراري: {str(e)}")
            raise

    def generate_thermal_receipt_text(self, invoice_data, items_data):
        """توليد نص الإيصال الحراري"""
        width = 32  # عرض الإيصال بالأحرف

        receipt = []
        receipt.append("=" * width)
        receipt.append("برنامج ست الكل للمحاسبة".center(width))
        receipt.append("فاتورة بيع".center(width))
        receipt.append("=" * width)
        receipt.append("")

        # معلومات الفاتورة
        receipt.append(f"رقم الفاتورة: {invoice_data['invoice_number']}")
        receipt.append(f"التاريخ: {invoice_data['invoice_date']}")
        receipt.append(f"العميل: {invoice_data.get('customer_name', 'عميل نقدي')}")
        receipt.append("-" * width)

        # الأصناف
        total_amount = 0
        for item in items_data:
            name = item.get('product_name', '')[:20]  # قطع الاسم إذا كان طويلاً
            qty = item.get('quantity', 0)
            price = item.get('unit_price', 0)
            line_total = item.get('line_total', 0)

            receipt.append(f"{name}")
            receipt.append(f"{qty:.1f} x {price:.2f} = {line_total:.2f}")
            receipt.append("")
            total_amount += line_total

        receipt.append("-" * width)

        # الإجماليات
        subtotal = invoice_data.get('subtotal', total_amount)
        discount = invoice_data.get('discount_amount', 0)
        tax = invoice_data.get('tax_amount', 0)
        final_total = invoice_data.get('total_amount', total_amount)

        receipt.append(f"المجموع الفرعي: {subtotal:.2f} ر.س".rjust(width))
        if discount > 0:
            receipt.append(f"الخصم: {discount:.2f} ر.س".rjust(width))
        if tax > 0:
            receipt.append(f"الضريبة: {tax:.2f} ر.س".rjust(width))
        receipt.append("=" * width)
        receipt.append(f"الإجمالي: {final_total:.2f} ر.س".rjust(width))
        receipt.append("=" * width)
        receipt.append("")
        receipt.append("شكراً لتعاملكم معنا".center(width))
        receipt.append("")

        return "\n".join(receipt)


def main():
    """اختبار نظام الطباعة"""
    printer = InvoicePrinter()

    # بيانات تجريبية
    invoice_data = {
        'invoice_number': 'SAL-2024-07-00001',
        'invoice_date': '2024-07-19',
        'customer_name': 'عميل تجريبي',
        'payment_method': 'نقدي',
        'subtotal': 100.00,
        'discount_amount': 5.00,
        'tax_amount': 15.00,
        'total_amount': 110.00,
        'notes': 'فاتورة تجريبية'
    }

    items_data = [
        {
            'product_code': 'P001',
            'product_name': 'منتج تجريبي 1',
            'quantity': 2,
            'unit_price': 25.00,
            'line_total': 50.00
        },
        {
            'product_code': 'P002',
            'product_name': 'منتج تجريبي 2',
            'quantity': 1,
            'unit_price': 50.00,
            'line_total': 50.00
        }
    ]

    try:
        # اختبار PDF
        if REPORTLAB_AVAILABLE:
            pdf_path = printer.print_invoice_pdf(invoice_data, items_data)
            print(f"تم إنشاء PDF: {pdf_path}")

        # اختبار Excel
        if OPENPYXL_AVAILABLE:
            excel_path = printer.export_invoice_excel(invoice_data, items_data)
            print(f"تم إنشاء Excel: {excel_path}")

        # اختبار الإيصال الحراري
        thermal_path = printer.print_thermal_receipt(invoice_data, items_data)
        print(f"تم إنشاء إيصال حراري: {thermal_path}")

    except Exception as e:
        print(f"خطأ في الاختبار: {str(e)}")


if __name__ == "__main__":
    main()
